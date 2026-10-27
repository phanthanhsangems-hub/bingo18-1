"""P277: GUI BU bao cao cuoi ngay cho mot ngay da qua (?ngay=YYYY-MM-DD).

23:55 hom 27/10 Cloud Run khong ket noi duoc Supabase -> bao cao mat han, va
endpoint chi biet tinh "hom nay" theo NOW() nen khong co cach nao gui lai.
Chay DUNG endpoint that tren SQLite, chan Telegram de bat noi dung.
"""
import io, json, os, sqlite3, sys, tempfile
from datetime import datetime, timedelta

os.environ.pop('DATABASE_URL', None)
os.environ['TELEGRAM_BOT_TOKEN'] = 'x'
os.environ['TELEGRAM_CHAT_ID'] = 'x'
_TMP = os.path.join(tempfile.mkdtemp(), 'bu.db')
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import config
config.DB_PATH = _TMP
import app as A
import telegram_bot as TB
import openpyxl

DAT = HONG = 0
def kiem(t, ok, ct=''):
    global DAT, HONG
    if ok: DAT += 1
    else: HONG += 1
    print(f"  {'DAT ' if ok else 'HONG'} {t}" + ('' if ok else f" -> {ct}"))

bat = {'msg': [], 'doc': []}
TB.TelegramBot.send_message = lambda self, m, *a, **k: bat['msg'].append(m) or True
TB.TelegramBot.send_document = lambda self, b, fn, cap='': bat['doc'].append((b, fn)) or True

vn_nay = datetime.utcnow() + timedelta(hours=7)
NGAY_CU = (vn_nay - timedelta(days=3)).strftime('%Y-%m-%d')
NGAY_GIUA = (vn_nay - timedelta(days=2)).strftime('%Y-%m-%d')

def nap(ngay_vn: str, dn0: int, so_ky: int, thang: int):
    goc_utc = datetime.strptime(ngay_vn, '%Y-%m-%d').replace(hour=12) - timedelta(hours=7)
    t = goc_utc.strftime('%Y-%m-%d %H:%M:%S')
    c = sqlite3.connect(_TMP)
    for k in range(so_ky):
        dn = dn0 + k
        c.execute("INSERT INTO draw_history (draw_number, draw_time, numbers, size_category, sum_value)"
                  " VALUES (?,?,?,?,?)", (dn, t, json.dumps([1, 2, 4]), 'NHO', 7))
        c.execute("INSERT INTO predictions (draw_number, model_name, predicted_numbers, confidence,"
                  " win_prob, prediction_time) VALUES (?,?,?,?,?,?)",
                  (dn, 'test', json.dumps([1, 2, 4]), 0.6, 0.39, t))
        pid = c.execute("SELECT id FROM predictions WHERE draw_number=?", (dn,)).fetchone()[0]
        c.execute("INSERT INTO prediction_results (prediction_id, draw_number, actual_numbers,"
                  " match_count, is_win, is_win_size, created_at) VALUES (?,?,?,?,?,?,?)",
                  (pid, dn, json.dumps([1, 2, 4]), 3, k < thang, k < thang, t))
    c.commit(); c.close()

A.app.test_client().get('/api/health')   # tao bang
nap(NGAY_CU, 700000, 12, 5)              # ngay can gui bu: 12 ky, thang 5
nap(vn_nay.strftime('%Y-%m-%d'), 800000, 7, 1)   # hom nay: 7 ky — khong duoc lot vao

c = A.app.test_client()
SEC = {'X-Trigger-Secret': 'bi-mat'}
A.config.TRIGGER_SECRET = 'bi-mat'

print("=== 1. BAO VE ===")
r = c.post(f'/api/daily-summary?ngay={NGAY_CU}')
kiem("khong secret -> 403", r.status_code == 403, r.status_code)
r = c.post("/api/daily-summary?ngay=2026-10-27'--", headers=SEC)
kiem("ngay sai dang -> 400 (khong cho chen SQL)", r.status_code == 400, r.status_code)
kiem("chua gui gi", not bat['msg'])

print("=== 2. GUI BU DUNG NGAY ===")
r = c.post(f'/api/daily-summary?ngay={NGAY_CU}&bu=1', headers=SEC)
d = r.get_json()
kiem("200", r.status_code == 200, (r.status_code, r.data[:200]))
kiem("dem dung 12 ky cua ngay do (khong lan hom nay)", d.get('total_predictions') == 12, d)
kiem("thang 5", d.get('wins') == 5, d)
m = bat['msg'][-1] if bat['msg'] else ''
dd = datetime.strptime(NGAY_CU, '%Y-%m-%d').strftime('%d/%m/%Y')
kiem("tieu de la ngay gui bu + '(gửi bù)'", f"TỔNG KẾT {dd}</b> (gửi bù)" in m, m[:120])
if bat['doc']:
    wb = openpyxl.load_workbook(io.BytesIO(bat['doc'][-1][0]))
    kiem("Excel 12 dong ky cua ngay do", wb.worksheets[0].max_row == 13, wb.worksheets[0].max_row)
    kiem("ten file theo ngay gui bu", bat['doc'][-1][1] == f"Bingo18_{NGAY_CU}.xlsx", bat['doc'][-1][1])
else:
    kiem("co gui Excel", False)

print("=== 3. KHONG GUI TRUNG ===")
n = len(bat['msg'])
r = c.post(f'/api/daily-summary?ngay={NGAY_CU}&bu=1', headers=SEC)
kiem("lan 2 -> da_gui", r.get_json().get('status') == 'da_gui', r.get_json())
kiem("khong gui them tin", len(bat['msg']) == n)
r = c.post(f'/api/daily-summary?ngay={NGAY_CU}', headers=SEC)
kiem("khong co bu=1 thi van gui (goi tay)", len(bat['msg']) == n + 1)

print("=== 4. MOC DA GUI CHI TIEN LEN ===")
r = c.post('/api/daily-summary', headers=SEC)          # 23:55 binh thuong: hom nay
kiem("bao cao hom nay 7 ky", r.get_json().get('total_predictions') == 7, r.get_json())
con = sqlite3.connect(_TMP)
moc = lambda: con.execute("SELECT config_value FROM system_config WHERE config_key='bao_cao_cuoi_ngay_da_gui'").fetchone()[0]
kiem("moc = hom nay", moc() == vn_nay.strftime('%Y-%m-%d'), moc())
nap(NGAY_GIUA, 750000, 3, 1)
c.post(f'/api/daily-summary?ngay={NGAY_GIUA}', headers=SEC)
kiem("gui bu ngay cu KHONG keo moc lui", moc() == vn_nay.strftime('%Y-%m-%d'), moc())

print("=== 5. 23:55 BINH THUONG KHONG DOI ===")
A.config.TRIGGER_SECRET = 'bi-mat'
r = c.post('/api/daily-summary')     # Cloud Scheduler khong gui secret — van chay
kiem("khong ?ngay, khong secret -> van 200", r.status_code == 200, r.status_code)

print(f"\nDAT: {DAT}   HONG: {HONG}")
sys.exit(1 if HONG else 0)
