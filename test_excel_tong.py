"""P219: sheet "Thống kê tổng" trong file Excel gửi cuối ngày.

Người dùng muốn biết trong ngày mỗi tổng 3..18 ra bao nhiêu lần.

Ba thứ dễ sai nhất, mỗi thứ một test:
  1. Số đếm phải khớp ĐÚNG dữ liệu — gom trong vòng lặp dựng sheet dữ liệu
     nên nếu vòng lặp đó đổi (bỏ qua dòng hỏng chẳng hạn) thì đếm lệch.
  2. Tổng các ô "Số lần" phải bằng đúng số kỳ trong ngày. Lệch = có kỳ bị
     đếm hai lần hoặc bị bỏ sót.
  3. Sheet dữ liệu CŨ không được đụng tới — thêm sheet mới không được làm
     xê dịch bảng kỳ quay mà người dùng vẫn đang dùng.

Chạy ĐÚNG endpoint /api/daily-summary thật, chặn Telegram lại để bắt file
Excel — không tự dựng lại workbook, vì như thế chỉ test code của chính test.
"""
import io
import json
import os
import sqlite3
import tempfile
from collections import Counter
from datetime import datetime, timedelta

os.environ.pop('DATABASE_URL', None)
os.environ['TELEGRAM_BOT_TOKEN'] = 'x'
os.environ['TELEGRAM_CHAT_ID']   = 'x'
_TMP = os.path.join(tempfile.mkdtemp(), 'xl.db')

import config
config.DB_PATH = _TMP
import app as A
import telegram_bot as TB
import openpyxl

fails = []
def ok(dk, ten, ct=''):
    print(f"  {'OK  ' if dk else 'HONG'}  {ten}")
    if not dk:
        fails.append(ten)
        if ct: print(f"        {ct}")

# ── Chặn Telegram, bắt lấy file ───────────────────────────────
_bat = {}
def _gui_doc(self, file_bytes, filename, caption=""):
    _bat['bytes'] = file_bytes; _bat['name'] = filename; _bat['caption'] = caption
    return True
def _gui_msg(self, *a, **k):
    return True
# app.py import TelegramBot BEN TRONG ham (app.py:282) nen phai va vao chinh
# module telegram_bot, khong phai thuoc tinh cua app.
TB.TelegramBot.send_document = _gui_doc
TB.TelegramBot.send_message  = _gui_msg

# ── Nạp dữ liệu: giờ VN = UTC + 7, nên lưu UTC lùi 7 tiếng ────
BO = {3:[1,1,1], 7:[1,2,4], 10:[2,4,4], 11:[4,4,3], 13:[3,4,6], 18:[6,6,6]}
KE_HOACH = {3: 2, 7: 20, 10: 25, 11: 18, 13: 30, 18: 1}   # tổng -> số kỳ

def nap():
    now_utc = datetime.utcnow()
    vn = now_utc + timedelta(hours=7)
    # neo vào 12:00 giờ VN hôm nay để chắc chắn không rơi sang ngày khác
    goc_vn  = vn.replace(hour=12, minute=0, second=0, microsecond=0)
    goc_utc = goc_vn - timedelta(hours=7)
    c = sqlite3.connect(_TMP)
    c.execute("DELETE FROM draw_history")
    c.execute("DELETE FROM prediction_results")
    c.execute("DELETE FROM predictions")
    dn = 500000
    for tong, sl in KE_HOACH.items():
        for _ in range(sl):
            dn += 1
            ns = BO[tong]
            sz = 'NHO' if tong <= 9 else ('HOA' if tong <= 11 else 'LON')
            c.execute("INSERT INTO draw_history (draw_number, draw_time, numbers,"
                      " size_category, sum_value) VALUES (?,?,?,?,?)",
                      (dn, goc_utc.strftime('%Y-%m-%d %H:%M:%S'), json.dumps(ns), sz, tong))
    # Endpoint dung som voi status 'empty' neu hom nay chua co prediction nao
    # (app.py:6042), tuc khoi Excel khong bao gio chay. Phai nap ca prediction
    # + ket qua thi moi di toi duoc doan dung sheet.
    for k in range(10):
        c.execute("INSERT INTO predictions (draw_number, model_name, predicted_numbers,"
                  " confidence, win_prob, prediction_time) VALUES (?,?,?,?,?,?)",
                  (500001 + k, 'test', json.dumps([1, 2, 4]), 0.6, 0.39,
                   goc_utc.strftime('%Y-%m-%d %H:%M:%S')))
        pid = c.execute("SELECT id FROM predictions WHERE draw_number=?",
                        (500001 + k,)).fetchone()[0]
        c.execute("INSERT INTO prediction_results (prediction_id, draw_number,"
                  " actual_numbers, match_count, is_win, is_win_size, created_at)"
                  " VALUES (?,?,?,?,?,?,?)",
                  (pid, 500001 + k, json.dumps([1, 2, 4]), 3, k < 4, k < 4,
                   goc_utc.strftime('%Y-%m-%d %H:%M:%S')))

    # kỳ của HÔM QUA — không được lọt vào báo cáo
    hq = goc_utc - timedelta(days=1)
    for k in range(5):
        dn += 1
        c.execute("INSERT INTO draw_history (draw_number, draw_time, numbers,"
                  " size_category, sum_value) VALUES (?,?,?,?,?)",
                  (dn, hq.strftime('%Y-%m-%d %H:%M:%S'), json.dumps([6,6,6]), 'LON', 18))
    c.commit(); c.close()

print("=" * 62)
print("P219: sheet Thong ke tong trong Excel cuoi ngay")
print("=" * 62)

nap()
r = A.app.test_client().post('/api/daily-summary',
                             headers={'X-Trigger-Secret': os.environ.get('TRIGGER_SECRET', '')})
ok(r.status_code == 200, "endpoint tra 200", f"status={r.status_code} body={r.data[:200]}")
ok('bytes' in _bat, "co gui file Excel")
if 'bytes' not in _bat:
    print("=" * 62); print("Khong bat duoc file — dung."); raise SystemExit(1)

wb = openpyxl.load_workbook(io.BytesIO(_bat['bytes']))
ok('Thống kê tổng' in wb.sheetnames, "co sheet 'Thong ke tong'", str(wb.sheetnames))
ok(len(wb.sheetnames) == 2, "dung 2 sheet", str(wb.sheetnames))

# ── Test 3: sheet dữ liệu CŨ không bị đụng ────────────────────
ws1 = wb.worksheets[0]
ok(ws1.title.startswith('Bingo18_'), "sheet du lieu van dung dau", ws1.title)
hdr = [c.value for c in ws1[1]]
ok(hdr == ["Kỳ", "Giờ VN", "Số 1", "Số 2", "Số 3", "Tổng", "Size"],
   "header sheet du lieu KHONG doi", str(hdr))
n_ky = sum(KE_HOACH.values())
ok(ws1.max_row == n_ky + 1, f"sheet du lieu van du {n_ky} dong", f"max_row={ws1.max_row}")

# ── Test 1: số đếm khớp đúng dữ liệu ──────────────────────────
ws2 = wb['Thống kê tổng']
doc = {}
for i in range(2, 18):
    t   = ws2.cell(row=i, column=1).value
    lan = ws2.cell(row=i, column=2).value
    doc[t] = lan
ok(sorted(doc) == list(range(3, 19)), "du 16 dong tong 3..18", str(sorted(doc)))
sai = {t: (doc.get(t), KE_HOACH.get(t, 0)) for t in range(3, 19)
       if doc.get(t) != KE_HOACH.get(t, 0)}
ok(not sai, "so lan KHOP dung du lieu nap vao", str(sai))
ok(doc[18] == 1, "ky HOM QUA khong lot vao (tong 18 chi 1 lan)", f"={doc[18]}, neu 6 la lot")

# ── Test 2: cộng lại phải bằng số kỳ trong ngày ───────────────
ok(sum(doc.values()) == n_ky,
   f"cong tat ca = {n_ky} ky trong ngay", f"={sum(doc.values())}")
o_cong = ws2.cell(row=19, column=2).value
ok(o_cong == n_ky, "dong CONG in dung so ky", f"={o_cong}")

# ── Kỳ vọng tính đúng ─────────────────────────────────────────
kv7 = ws2.cell(row=2 + (7 - 3), column=4).value
mong = round(n_ky * 15 / 216.0, 1)
ok(abs(kv7 - mong) < 0.05, f"ky vong tong 7 = {mong}", f"={kv7}")
chenh7 = ws2.cell(row=2 + (7 - 3), column=5).value
ok(abs(chenh7 - round(KE_HOACH[7] - n_ky * 15 / 216.0, 1)) < 0.05,
   "chenh = so lan - ky vong", f"={chenh7}")

# ── Gộp theo SIZE ─────────────────────────────────────────────
nho = sum(v for t, v in KE_HOACH.items() if t <= 9)
hoa = sum(v for t, v in KE_HOACH.items() if 10 <= t <= 11)
lon = sum(v for t, v in KE_HOACH.items() if t >= 12)
ok(ws2.cell(row=22, column=2).value == nho, f"NHO = {nho}", str(ws2.cell(row=22, column=2).value))
ok(ws2.cell(row=23, column=2).value == hoa, f"HOA = {hoa}", str(ws2.cell(row=23, column=2).value))
ok(ws2.cell(row=24, column=2).value == lon, f"LON = {lon}", str(ws2.cell(row=24, column=2).value))
ok(nho + hoa + lon == n_ky, "ba SIZE cong lai = tong so ky")

# ── Chú thích có mặt ──────────────────────────────────────────
ok('Kỳ vọng' in str(ws2.cell(row=26, column=1).value or ''),
   "co dong giai thich ky vong")

print("=" * 62)
if fails:
    print(f"HONG {len(fails)} test:")
    for f in fails: print("  - " + f)
    raise SystemExit(1)
print("Tat ca test P219 deu dat.")
